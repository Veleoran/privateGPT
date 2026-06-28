# -*- coding: utf-8 -*-
"""CR Fibre v4 — export Excel des tickets sur une période, avec total de points."""
import io
import datetime as dt

import db


def current_week():
    today = dt.date.today()
    start = today - dt.timedelta(days=today.weekday())  # lundi
    end = start + dt.timedelta(days=6)                   # dimanche
    return start.isoformat(), end.isoformat()


def current_month():
    today = dt.date.today()
    start = today.replace(day=1)
    if start.month == 12:
        nxt = start.replace(year=start.year + 1, month=1)
    else:
        nxt = start.replace(month=start.month + 1)
    end = nxt - dt.timedelta(days=1)
    return start.isoformat(), end.isoformat()


def build_xlsx(start, end):
    """Construit un classeur xlsx en mémoire. Retourne (BytesIO, filename)."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

    rows = db.list_period(start, end)

    wb = Workbook()
    ws = wb.active
    ws.title = "Tickets"

    headers = ["Date", "Type", "Équipe", "Effectif", "Statut", "Mail", "Points"]
    ws.append(headers)

    head_fill = PatternFill("solid", fgColor="1F2A44")
    head_font = Font(bold=True, color="FFFFFF")
    thin = Side(style="thin", color="D0D0D0")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    for c in ws[1]:
        c.fill = head_fill
        c.font = head_font
        c.alignment = Alignment(horizontal="center")
        c.border = border

    total = 0.0
    for r in rows:
        date = r["date_intervention"] or (r["created_at"] or "")[:10]
        try:
            d, m, y = date.split("-")[2], date.split("-")[1], date.split("-")[0]
            date = f"{d}/{m}/{y}"
        except Exception:
            pass
        pts = r["points"] or 0
        total += pts
        ws.append([
            date,
            r["template_label"] or r["template_key"],
            r["equipe"] or "",
            r["effectif"] or "",
            r["statut"] or "",
            "Oui" if r["mail_envoye"] else "Non",
            pts,
        ])

    # Ligne total
    ws.append([])
    last = ws.max_row + 1
    ws.append(["", "", "", "", "", "TOTAL", round(total, 2)])
    for c in ws[ws.max_row]:
        c.font = Font(bold=True)
    ws.cell(row=ws.max_row, column=6).alignment = Alignment(horizontal="right")

    # Largeurs
    widths = [12, 30, 22, 10, 22, 8, 9]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[chr(64 + i)].width = w
    ws.freeze_panes = "A2"

    bio = io.BytesIO()
    wb.save(bio)
    bio.seek(0)
    fname = f"cr-fibre_{start}_au_{end}.xlsx"
    return bio, fname
