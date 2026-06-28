# -*- coding: utf-8 -*-
"""CR Fibre v4.2 — export Excel mensuel de paie.

Deux tableaux (feuilles) :
  • « À facturer » : interventions Clos / Clos+Travaux, avec colonne Code et total points.
  • « Travaux »    : interventions Travaux / Clos+Travaux, avec DPR (suivi des reprises).
"""
import io
import datetime as dt

import db


def current_week():
    today = dt.date.today()
    start = today - dt.timedelta(days=today.weekday())  # lundi
    end = start + dt.timedelta(days=6)                   # dimanche
    return start.isoformat(), end.isoformat()


def current_month():
    today = dt.date.today()
    start = today.replace(day=1)
    if start.month == 12:
        nxt = start.replace(year=start.year + 1, month=1)
    else:
        nxt = start.replace(month=start.month + 1)
    end = nxt - dt.timedelta(days=1)
    return start.isoformat(), end.isoformat()


def _fr(date):
    try:
        p = date.split("-")
        return f"{p[2]}/{p[1]}/{p[0]}"
    except Exception:
        return date or ""


def _row_date(r):
    return r["date_intervention"] or (r["created_at"] or "")[:10]


def build_xlsx(start, end):
    """Construit un classeur xlsx (2 feuilles) en mémoire. Retourne (BytesIO, filename)."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

    rows = db.list_period(start, end)

    head_fill = PatternFill("solid", fgColor="1F2A44")
    head_font = Font(bold=True, color="FFFFFF")
    thin = Side(style="thin", color="D0D0D0")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    def style_header(ws):
        for c in ws[1]:
            c.fill = head_fill
            c.font = head_font
            c.alignment = Alignment(horizontal="center")
            c.border = border

    def set_widths(ws, widths):
        for i, w in enumerate(widths, start=1):
            ws.column_dimensions[chr(64 + i)].width = w
        ws.freeze_panes = "A2"

    def st(r):
        return (r["statut"] or "")

    billable = [r for r in rows if st(r) in ("Clos", "Clos/Travaux")]
    travaux = [r for r in rows if "Travaux" in st(r)]  # Travaux + Clos/Travaux

    wb = Workbook()

    # --- Feuille 1 : À facturer ---------------------------------------------
    ws = wb.active
    ws.title = "À facturer"
    ws.append(["Date", "Code", "Type", "Statut", "Effectif", "Points"])
    style_header(ws)
    total = 0.0
    for r in billable:
        pts = r["points"] or 0
        total += pts
        ws.append([
            _fr(_row_date(r)),
            r["code"] or "",
            r["template_label"] or r["template_key"],
            st(r),
            r["effectif"] or "",
            pts,
        ])
    ws.append([])
    ws.append(["", "", "", "", "TOTAL", round(total, 2)])
    for c in ws[ws.max_row]:
        c.font = Font(bold=True)
    ws.cell(row=ws.max_row, column=6).alignment = Alignment(horizontal="right")
    set_widths(ws, [12, 16, 34, 14, 10, 9])

    # --- Feuille 2 : Travaux (suivi reprises) -------------------------------
    wt = wb.create_sheet("Travaux")
    wt.append(["Date", "Code", "Type", "Statut", "DPR"])
    style_header(wt)
    for r in travaux:
        wt.append([
            _fr(_row_date(r)),
            r["code"] or "",
            r["template_label"] or r["template_key"],
            st(r),
            _fr(r["dpr"] or ""),
        ])
    set_widths(wt, [12, 16, 34, 14, 12])

    bio = io.BytesIO()
    wb.save(bio)
    bio.seek(0)
    fname = f"cr-fibre_{start}_au_{end}.xlsx"
    return bio, fname
